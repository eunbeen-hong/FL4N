"""Build figure_ppts/teaser.pptx (paper teaser: explicit Gaussians vs FL4N latent) and a PNG preview via an HTML twin of the same layout; crops come from the CTVS prediction folders."""
import os, re, json, random, subprocess
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from lxml import etree

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
B = '/home/cvlab20/project/eunbeen/ssd_symlink/sde/eunbeen/baseline_evals_res/ctvs/nvidia/nvidia'
B_MOVIES = '/home/cvlab20/project/eunbeen/ssd_symlink/sde/eunbeen/final_benchmark/ctvs/nvidia/nvidia'
OUT = os.path.join(ROOT, 'figure_ppts', 'teaser.pptx')
IMG = os.path.join(ROOT, 'tools', 'teaser_ppt', 'img')
PW = os.environ.get('PW', '/tmp/claude-1001/-home-cvlab20-project-eunbeen-LagerNVS4d/2029518d-9a2f-403f-a387-494dcea53f53/scratchpad/pw')
FONT = 'Google Sans'
C = dict(ink='231a12', ink2='4e4033', ink3='7f6d5c', line='efe3d2', line2='e2d0b8', p='b4591e', pd='8e4113', ps='f4d9c2', pg='fbf1e6', y='e8b53a', ys='fbe6a8', g='1b7f4b', r='c0392b', slate='e6ecef', slate2='a3b4bd', gray='f3efe9', white='ffffff')
W, H = 13.333, 6.2
EX = [('Balloon1', 'movies', 4, 10, (330, 20, 830, 395), 'tearing'), ('Jumping', '4dgt', 0, 8, (80, 70, 580, 445), 'blur'), ('Playground', 'neoverse_tto_nostatic', 2, 10, (340, 60, 800, 405), 'holes')]


def pred(scene, key, v, f):
    d = f'{B_MOVIES if key == "movies" else B}/{key}/{scene}/pred'
    p = f'{d}/v{v:02d}_f{f:02d}.png'
    if not os.path.exists(p):
        p = [f'{d}/{x}' for x in os.listdir(d) if re.match(rf'^all_{v}_i\d+_f{f}\.png$', x)][0]
    return Image.open(p).convert('RGB')


def crops():
    os.makedirs(IMG, exist_ok=True)
    out = []
    for scene, key, v, f, box, chip in EX:
        a, b = f'{IMG}/{scene}_base.png', f'{IMG}/{scene}_ours.png'
        pred(scene, key, v, f).crop(box).save(a)
        pred(scene, 'ours_posed', v, f).crop(box).save(b)
        gt = f'{IMG}/{scene}_in.png'
        Image.open(f'/mnt/sde/eunbeen/data/NVIDIA_x/NVIDIA/{scene}/multiview_GT/{f + 1:08d}/cam{v + 1:02d}.jpg').convert('RGB').save(gt)
        out.append((a, b, gt, chip))
    logo = f'{IMG}/flan.png'
    Image.open(f'{ROOT}/figure_ppts/flan.png').save(logo)
    return out, logo


def layout(ex, logo):
    """Returns a list of primitives: (kind, dict). Units are inches; kinds: rect, text, image, line, ellipse."""
    L = []
    def rect(x, y, w, h, fill=None, stroke=None, sw=1, r=0.12, alpha=None, name=''):
        L.append(('rect', dict(x=x, y=y, w=w, h=h, fill=fill, stroke=stroke, sw=sw, r=r, alpha=alpha, name=name)))
    def text(x, y, w, h, s, size, color='ink', bold=False, align='l', anchor='m'):
        L.append(('text', dict(x=x, y=y, w=w, h=h, s=s, size=size, color=color, bold=bold, align=align, anchor=anchor)))
    def image(x, y, w, h, path, stroke=None, sw=1.5):
        L.append(('image', dict(x=x, y=y, w=w, h=h, path=path, stroke=stroke, sw=sw)))
    def line(x1, y1, x2, y2, color='ink', sw=1.75, arrow=False):
        L.append(('line', dict(x1=x1, y1=y1, x2=x2, y2=y2, color=color, sw=sw, arrow=arrow)))
    def ell(x, y, w, h, fill, rot=0, alpha=None, stroke=None):
        L.append(('ellipse', dict(x=x, y=y, w=w, h=h, fill=fill, rot=rot, alpha=alpha, stroke=stroke)))
    image(0.4, 0.22, 0.9, 0.9, logo)
    text(1.35, 0.2, 2.2, 0.95, 'FL4N', 38, 'p', True)
    text(3.35, 0.2, 7, 0.95, 'Novel views and timestamps from one monocular video', 17, 'ink2', False)
    rows = [dict(y=1.35, h=2.2, fill='gray', stroke='line2', tag='Explicit  ·  3D Gaussians', tagc='ink2', tagbg='white'), dict(y=3.75, h=2.2, fill='pg', stroke='ps', tag='Latent  ·  FL4N (ours)', tagc='pd', tagbg='white')]
    mid = (rows[0]['y'] + rows[1]['y'] + rows[1]['h']) / 2
    sw, sh = 1.7, 0.96
    for i, (a, b, gt, chip) in enumerate(reversed(ex)):
        k = 2 - i
        image(0.45 + k * 0.18, mid - sh / 2 - 0.5 + k * 0.28, sw, sh, gt, stroke='p' if k == 0 else 'line2', sw=1.5)
    text(0.4, mid + 0.75, 2.5, 0.35, 'Monocular video', 13, 'ink3', False, 'c')
    bx = 0.45 + 2 * 0.18 + sw + 0.1
    line(bx, mid, bx + 0.3, mid)
    line(bx + 0.3, rows[0]['y'] + rows[0]['h'] / 2, bx + 0.3, rows[1]['y'] + rows[1]['h'] / 2, sw=1.75)
    for r in rows:
        line(bx + 0.3, r['y'] + r['h'] / 2, 3.15, r['y'] + r['h'] / 2, arrow=True)
    for ri, r in enumerate(rows):
        rect(3.2, r['y'], W - 3.2 - 0.35, r['h'], fill=r['fill'], stroke=r['stroke'], sw=1, r=0.18)
        rect(3.45, r['y'] - 0.2, 2.6 if ri == 0 else 2.3, 0.4, fill=r['tagbg'], stroke=r['stroke'], sw=1, r=0.2)
        text(3.45, r['y'] - 0.2, 2.6 if ri == 0 else 2.3, 0.4, r['tag'], 13, r['tagc'], True, 'c')
        cy = r['y'] + r['h'] / 2 + 0.08
        ix = 3.5
        if ri == 0:
            random.seed(7)
            for n in range(34):
                ang = random.random() * 6.283; rad = random.random() ** .5 * .62
                cx, cyy = ix + .85 + rad * 1.05 * __import__('math').cos(ang), cy + rad * .82 * __import__('math').sin(ang)
                w, h = random.uniform(.16, .42), random.uniform(.09, .2)
                ell(cx - w / 2, cyy - h / 2, w, h, random.choice(['slate2', 'ink3', 'line2', 'p', 'y', 'ink2']), rot=random.uniform(0, 180), alpha=random.uniform(.35, .7))
            text(ix + .1, cy + .68, 1.5, .3, '3D Gaussians', 11, 'ink3', False, 'c')
        else:
            for n in range(2, -1, -1):
                ox, oy = ix + .35 + n * .11, cy - .55 - n * .11 + .25
                for gx in range(4):
                    for gy in range(4):
                        hot = n == 0 and (gx, gy) in [(2, 0), (1, 1), (2, 1), (2, 2), (3, 2)]
                        rect(ox + gx * .2, oy + gy * .2, .2, .2, fill='ys' if hot else ('ps' if n == 0 else 'slate'), stroke='p' if n == 0 else 'slate2', sw=.75, r=0)
            text(ix + .1, cy + .68, 1.5, .3, '4D latent', 11, 'pd', False, 'c')
        line(ix + 1.75, cy, ix + 2.2, cy, arrow=True)
        iw, ih = 2.14, 1.605
        for n, (a, b, gt, chip) in enumerate(ex):
            x = ix + 2.35 + n * (iw + 0.22)
            y = cy - ih / 2 - 0.06
            image(x, y, iw, ih, a if ri == 0 else b, stroke='line2' if ri == 0 else 'p', sw=1.5)
            if ri == 0:
                rect(x + iw - 0.95, y + ih - 0.42, 0.85, 0.32, fill='ys', stroke=None, r=0.1)
                text(x + iw - 0.95, y + ih - 0.42, 0.85, 0.32, chip, 10.5, 'ink', True, 'c')
                ell(x + 0.1, y + 0.1, 0.32, 0.32, 'r')
                text(x + 0.1, y + 0.1, 0.32, 0.32, '✕', 11, 'white', True, 'c')
            else:
                ell(x + 0.1, y + 0.1, 0.32, 0.32, 'g')
                text(x + 0.1, y + 0.1, 0.32, 0.32, '✓', 11, 'white', True, 'c')
    return L


def alpha(fill_elm, a):
    clr = fill_elm.find(qn('a:srgbClr'))
    e = etree.SubElement(clr, qn('a:alpha')); e.set('val', str(int(a * 100000)))


def emit_pptx(L, path):
    prs = Presentation(); prs.slide_width, prs.slide_height = Inches(W), Inches(H)
    s = prs.slides.add_slide(prs.slide_layouts[6])
    for kind, d in L:
        if kind in ('rect', 'ellipse'):
            shp = s.shapes.add_shape(MSO_SHAPE.OVAL if kind == 'ellipse' else (MSO_SHAPE.ROUNDED_RECTANGLE if d.get('r') else MSO_SHAPE.RECTANGLE), Inches(d['x']), Inches(d['y']), Inches(d['w']), Inches(d['h']))
            if kind == 'rect' and d.get('r'):
                shp.adjustments[0] = min(0.5, d['r'] / min(d['w'], d['h']))
            if d.get('fill'):
                shp.fill.solid(); shp.fill.fore_color.rgb = RGBColor.from_string(C[d['fill']])
                if d.get('alpha'):
                    alpha(shp.fill._xPr.find(qn('a:solidFill')), d['alpha'])
            else:
                shp.fill.background()
            if d.get('stroke'):
                shp.line.color.rgb = RGBColor.from_string(C[d['stroke']]); shp.line.width = Pt(d.get('sw', 1))
            else:
                shp.line.fill.background()
            shp.shadow.inherit = False
            if d.get('rot'):
                shp.rotation = d['rot']
            shp.text_frame.text = ''
        elif kind == 'text':
            tb = s.shapes.add_textbox(Inches(d['x']), Inches(d['y']), Inches(d['w']), Inches(d['h']))
            tf = tb.text_frame; tf.word_wrap = True; tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
            tf.vertical_anchor = {'m': MSO_ANCHOR.MIDDLE, 't': MSO_ANCHOR.TOP}[d['anchor']]
            p = tf.paragraphs[0]; p.alignment = {'l': PP_ALIGN.LEFT, 'c': PP_ALIGN.CENTER, 'r': PP_ALIGN.RIGHT}[d['align']]
            run = p.add_run(); run.text = d['s']; run.font.size = Pt(d['size']); run.font.bold = d['bold']; run.font.name = FONT; run.font.color.rgb = RGBColor.from_string(C[d['color']])
        elif kind == 'image':
            pic = s.shapes.add_picture(d['path'], Inches(d['x']), Inches(d['y']), Inches(d['w']), Inches(d['h']))
            if d.get('stroke'):
                pic.line.color.rgb = RGBColor.from_string(C[d['stroke']]); pic.line.width = Pt(d['sw'])
        elif kind == 'line':
            c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(d['x1']), Inches(d['y1']), Inches(d['x2']), Inches(d['y2']))
            c.line.color.rgb = RGBColor.from_string(C[d['color']]); c.line.width = Pt(d['sw'])
            if d['arrow']:
                ln = c.line._get_or_add_ln(); te = etree.SubElement(ln, qn('a:tailEnd')); te.set('type', 'triangle'); te.set('w', 'med'); te.set('len', 'med')
    prs.save(path)


def emit_html(L, path):
    px = lambda v: f'{v * 120:.1f}px'
    h = [f'<!doctype html><html><head><meta charset="utf-8"><link href="https://fonts.googleapis.com/css2?family=Google+Sans:wght@400;500;700&display=swap" rel="stylesheet"><style>body{{margin:0;background:#fff;font-family:"Google Sans",Inter,sans-serif}}#s{{position:relative;width:{px(W)};height:{px(H)};overflow:hidden;background:#fff}}#s>*{{position:absolute;box-sizing:border-box}}</style></head><body><div id="s">']
    for kind, d in L:
        if kind in ('rect', 'ellipse'):
            st = f'left:{px(d["x"])};top:{px(d["y"])};width:{px(d["w"])};height:{px(d["h"])};'
            if d.get('fill'):
                c = C[d['fill']]; a = d.get('alpha')
                st += f'background:#{c}{int(a * 255):02x};' if a else f'background:#{c};'
            if d.get('stroke'):
                st += f'border:{d.get("sw", 1) * 1.33:.1f}px solid #{C[d["stroke"]]};'
            st += 'border-radius:50%;' if kind == 'ellipse' else f'border-radius:{px(d.get("r", 0))};'
            if d.get('rot'):
                st += f'transform:rotate({d["rot"]:.1f}deg);'
            h.append(f'<div style="{st}"></div>')
        elif kind == 'text':
            al = {'l': 'flex-start', 'c': 'center', 'r': 'flex-end'}[d['align']]
            h.append(f'<div style="left:{px(d["x"])};top:{px(d["y"])};width:{px(d["w"])};height:{px(d["h"])};display:flex;align-items:{"center" if d["anchor"] == "m" else "flex-start"};justify-content:{al};font-size:{d["size"] * 1.667:.1f}px;font-weight:{700 if d["bold"] else 400};color:#{C[d["color"]]};line-height:1.15;text-align:{ {"l": "left", "c": "center", "r": "right"}[d["align"]] }">{d["s"]}</div>')
        elif kind == 'image':
            bd = f'border:{d["sw"] * 1.33:.1f}px solid #{C[d["stroke"]]};' if d.get('stroke') else ''
            h.append(f'<img src="file://{d["path"]}" style="left:{px(d["x"])};top:{px(d["y"])};width:{px(d["w"])};height:{px(d["h"])};object-fit:fill;{bd}">')
        elif kind == 'line':
            x1, y1, x2, y2 = d['x1'] * 120, d['y1'] * 120, d['x2'] * 120, d['y2'] * 120
            mk = '<defs><marker id="m" markerWidth="8" markerHeight="8" refX="7" refY="4" orient="auto" markerUnits="userSpaceOnUse"><path d="M0,0 L8,4 L0,8 z" fill="#%s"/></marker></defs>' % C[d['color']]
            h.append(f'<svg style="left:0;top:0;width:{px(W)};height:{px(H)};pointer-events:none;overflow:visible">{mk}<line x1="{x1:.1f}" y1="{y1:.1f}" x2="{x2:.1f}" y2="{y2:.1f}" stroke="#{C[d["color"]]}" stroke-width="{d["sw"] * 1.33:.1f}" {"marker-end=url(#m)" if d["arrow"] else ""}/></svg>')
    h.append('</div></body></html>')
    open(path, 'w').write('\n'.join(h))


def main():
    ex, logo = crops()
    L = layout(ex, logo)
    emit_pptx(L, OUT)
    html = os.path.join(IMG, 'preview.html'); emit_html(L, html)
    png = os.path.join(ROOT, 'tools', 'teaser_ppt', 'teaser_preview.png')
    js = os.path.join(IMG, 'shot.js')
    open(js, 'w').write("const { chromium } = require('playwright');(async () => { const b = await chromium.launch({ args: ['--allow-file-access-from-files'] }); const p = await b.newPage({ viewport: { width: %d, height: %d }, deviceScaleFactor: 2 }); await p.goto('file://%s'); await p.waitForTimeout(1500); await p.screenshot({ path: '%s' }); await b.close(); })();" % (round(W * 120), round(H * 120), html, png))
    subprocess.run(['node', js], cwd=PW, env=dict(os.environ, NODE_PATH=os.path.join(PW, 'node_modules')), check=True)
    print(OUT, os.path.getsize(OUT) // 1024, 'KB;', png)


if __name__ == '__main__':
    main()
