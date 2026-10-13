"""Extract pictures and layout from figure_ppts/*.pptx into static/figures/<fig>/ plus static/js/figdata.js describing each figure as rows/cols."""
import os, io, re, json, hashlib, sys
from pptx import Presentation
from PIL import Image

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, 'static', 'figures')
IN = 914400.0


def walk(shapes, ctx=(1, 1, 0, 0), depth=0):
    sx, sy, tx, ty = ctx
    for sh in shapes:
        if sh.shape_type == 6:
            x = sh._element.grpSpPr.find('{http://schemas.openxmlformats.org/drawingml/2006/main}xfrm')
            cho, che = x.chOff, x.chExt
            cw, ch = (che.cx or 1), (che.cy or 1)
            nsx, nsy = sx * x.ext.cx / cw, sy * x.ext.cy / ch
            ntx = tx + sx * x.off.x - nsx * cho.x
            nty = ty + sy * x.off.y - nsy * cho.y
            yield from walk(sh.shapes, (nsx, nsy, ntx, nty), depth + 1)
            continue
        x, y = (tx + sh.left * sx) / IN, (ty + sh.top * sy) / IN
        w, h = sh.width * sx / IN, sh.height * sy / IN
        yield sh, x, y, w, h


def save_img(sh, folder, maxdim=1024):
    blob = sh.image.blob
    key = hashlib.md5(blob).hexdigest()[:10]
    fn = f'{key}.webp'
    path = os.path.join(folder, fn)
    if not os.path.exists(path):
        im = Image.open(io.BytesIO(blob))
        if max(im.size) > maxdim:
            im.thumbnail((maxdim, maxdim), Image.LANCZOS)
        if im.mode not in ('RGB', 'RGBA'):
            im = im.convert('RGBA')
        im.save(path, 'WEBP', quality=86, method=6)
    return fn


def line_color(sh):
    try:
        ln = sh._element.spPr.find('{http://schemas.openxmlformats.org/drawingml/2006/main}ln')
        if ln is None:
            return None
        c = ln.find('.//{http://schemas.openxmlformats.org/drawingml/2006/main}srgbClr')
        return '#' + c.get('val').lower() if c is not None else None
    except Exception:
        return None


def grid(items, rows_by='y', tol=0.2):
    rows = []
    for it in sorted(items, key=lambda t: (t['y'], t['x'])):
        for r in rows:
            if abs(r[0]['y'] - it['y']) < tol:
                r.append(it)
                break
        else:
            rows.append([it])
    return [sorted(r, key=lambda t: t['x']) for r in rows]


def pics(pres, slide, folder, maxdim=1024, filt=None):
    out = []
    for sh, x, y, w, h in walk(pres.slides[slide].shapes):
        if sh.shape_type == 13 and (filt is None or filt(sh, x, y, w, h)):
            out.append(dict(x=round(x, 3), y=round(y, 3), w=round(w, 3), h=round(h, 3), img=save_img(sh, folder, maxdim), name=sh.name, ln=line_color(sh)))
    return out


def texts(pres, slide):
    out = []
    for sh, x, y, w, h in walk(pres.slides[slide].shapes):
        if sh.has_text_frame and sh.text_frame.text.strip():
            out.append(dict(x=round(x, 3), y=round(y, 3), w=round(w, 3), h=round(h, 3), t=sh.text_frame.text.strip()))
    return out


def fig(name):
    folder = os.path.join(OUT, name)
    os.makedirs(folder, exist_ok=True)
    return Presentation(os.path.join(ROOT, 'figure_ppts', f'{name}.pptx')), folder


def main():
    data = {}
    p, f = fig('dynfeature')
    rows = grid(pics(p, 0, f))
    data['dynfeature'] = dict(rows=[[c['img'] for c in r] for r in rows], labels=['Context frames', 'Static latent z_st', 'Residual latent z_res'])
    p, f = fig('gatefeature')
    rows = grid(pics(p, 0, f))
    data['gatefeature'] = dict(rows=[[c['img'] for c in r] for r in rows], labels=['Context frames', 'Gate activation'])
    p, f = fig('dynmask')
    rows = grid(pics(p, 0, f))
    data['dynmask'] = dict(rows=[[c['img'] for c in r] for r in rows], labels=['VGGT4D', 'FL4N (Ours)', 'Ground truth'])
    p, f = fig('supp_staticfeat')
    rows = grid(pics(p, 0, f))
    boxes = []
    for sh, x, y, w, h in walk(p.slides[0].shapes):
        if sh.shape_type == 1 and sh.name.startswith('Rectangle'):
            boxes.append(dict(x=x, y=y, w=w, h=h))
    bx = []
    for b in boxes:
        for ri, r in enumerate(rows):
            for ci, c in enumerate(r):
                if c['x'] - 0.02 <= b['x'] < c['x'] + c['w'] and c['y'] - 0.02 <= b['y'] < c['y'] + c['h']:
                    bx.append(dict(r=ri, c=ci, l=round((b['x'] - c['x']) / c['w'], 3), t=round((b['y'] - c['y']) / c['h'], 3), w=round(b['w'] / c['w'], 3), h=round(b['h'] / c['h'], 3)))
    data['staticfeat'] = dict(rows=[[c['img'] for c in r] for r in rows], labels=['Context frames', 'Static latent z_st'], boxes=bx)
    p, f = fig('supp_nn_residualfeat')
    allp = pics(p, 0, f)
    left = grid([c for c in allp if c['x'] < 9.1])
    right = grid([c for c in allp if c['x'] >= 9.1])
    data['residualgrid'] = dict(rgb=[[c['img'] for c in r] for r in left], res=[[c['img'] for c in r] for r in right])
    p, f = fig('evalsettings')
    allp = pics(p, 0, f, 512)
    panels = []
    for lo, hi, key, title in [(1.5, 4.4, 'nvs', 'Novel view synthesis'), (4.4, 8.0, 'ctvs', 'Cross-time view synthesis'), (8.0, 12, 'nts', 'Novel time synthesis')]:
        rows = grid([c for c in allp if lo <= c['x'] < hi], tol=0.15)
        panels.append(dict(key=key, title=title, rows=[[dict(img=c['img'], role='ctx' if (c['ln'] or '').lower() == '#ff0000' else ('tgt' if c['ln'] else 'none')) for c in r] for r in rows]))
    data['evalsettings'] = panels
    p, f = fig('compare_static')
    allp = pics(p, 0, f, 640)
    scenes = []
    for lo, hi, title in [(1.5, 6.1, 'Lady with a plate (CamxTime)'), (6.1, 11, 'Colonnade (Charge)')]:
        cs = [c for c in allp if lo <= c['x'] < hi and c['w'] < 2]
        inputs = sorted([c for c in cs if c['w'] < 0.55], key=lambda c: c['x'])
        big = [c for c in cs if c['w'] > 1.2][0]
        mids = grid([c for c in cs if 0.55 <= c['w'] <= 1.2])
        scenes.append(dict(title=title, inputs=[c['img'] for c in inputs], static=big['img'], gt=[c['img'] for c in mids[0]], ours=[c['img'] for c in mids[1]]))
    data['comparestatic'] = scenes
    p, f = fig('qual')
    allp = pics(p, 0, f, 672)
    tx = texts(p, 0)
    cols = sorted([t for t in tx if t['y'] < 11.5], key=lambda t: t['x'])
    rows = grid(allp, tol=0.5)
    rowlab = sorted([t for t in tx if t['y'] > 11.5], key=lambda t: t['y'])
    data['qual'] = dict(models=[t['t'] for t in cols], datasets=[t['t'] for t in rowlab], rows=[[c['img'] for c in r] for r in rows])
    json.dump(data, open(os.path.join(OUT, 'figdata.json'), 'w'), indent=0)
    open(os.path.join(ROOT, 'static', 'js', 'figdata.js'), 'w').write('window.FL4N_FIG = ' + json.dumps(data, separators=(',', ':')) + ';\n')
    for k, v in data.items():
        print(k, json.dumps(v)[:160])


if __name__ == '__main__':
    main()
